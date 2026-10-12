"""Build proposal.docx, proposal.tex/pdf and proposal.txt from docs/proposal.md.

``docs/proposal.md`` is the single source of truth for the written report. This
script renders it into the three formats the submission needs, so the Word file,
the PDF and the plain-text version can never drift apart.

There is no pandoc on this machine, so it implements the Markdown subset the
proposal actually uses -- headings, paragraphs, bullet and numbered lists,
pipe tables, block quotes, figures, and inline bold/italic/code. That subset is
checked by :func:`validate`, which fails loudly on syntax it would otherwise
render wrongly, rather than silently dropping it.
"""
from __future__ import annotations

import os
import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SRC = DOCS / "proposal.md"

XELATEX = Path.home() / "opt/texlive/2026/bin/x86_64-linux/xelatex"


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
@dataclass
class Block:
    kind: str                       # heading|para|bullets|numbers|table|figure|quote
    level: int = 0
    text: str = ""
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    caption: str = ""
    path: str = ""


def parse(md: str) -> list[Block]:
    lines = md.split("\n")
    blocks: list[Block] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("<!--"):                      # comment
            while i < len(lines) and "-->" not in lines[i]:
                i += 1
            i += 1
            continue

        if m := re.match(r"^(#{1,4})\s+(.*)$", stripped):
            blocks.append(Block("heading", level=len(m.group(1)), text=m.group(2).strip()))
            i += 1
            continue

        if m := re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped):
            blocks.append(Block("figure", caption=m.group(1), path=m.group(2)))
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", rows=rows))
            continue

        if re.match(r"^[-*]\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^[-*]\s+", lines[i].strip()):
                items.append(re.sub(r"^[-*]\s+", "", lines[i].strip()))
                i += 1
            blocks.append(Block("bullets", items=items))
            continue

        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while i < len(lines) and re.match(r"^\d+\.\s+", lines[i].strip()):
                items.append(re.sub(r"^\d+\.\s+", "", lines[i].strip()))
                i += 1
            blocks.append(Block("numbers", items=items))
            continue

        if stripped.startswith(">"):
            parts = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                parts.append(lines[i].strip().lstrip("> ").strip())
                i += 1
            blocks.append(Block("quote", text=" ".join(parts)))
            continue

        para = []
        while i < len(lines) and lines[i].strip() and not re.match(
            r"^(#{1,4}\s|[-*]\s|\d+\.\s|\||>|!\[)", lines[i].strip()
        ):
            para.append(lines[i].strip())
            i += 1
        blocks.append(Block("para", text=" ".join(para)))
    return blocks


def validate(md: str) -> None:
    """Refuse to build on Markdown features this converter does not render."""
    problems = []
    if re.search(r"^```", md, re.M):
        problems.append("fenced code blocks are not supported by this converter")
    if re.search(r"^\s{2,}[-*]\s", md, re.M):
        problems.append("nested lists are not supported")
    if re.search(r"\[([^\]]+)\]\((?!.*\.(png|jpg|pdf))", md):
        problems.append("hyperlinks are not supported; write the URL as plain text")
    if problems:
        raise SystemExit("docs/proposal.md uses unsupported syntax:\n  - "
                         + "\n  - ".join(problems))


# --------------------------------------------------------------------------- #
# Inline formatting
# --------------------------------------------------------------------------- #
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def split_inline(text: str) -> list[tuple[str, str]]:
    """Split a string into (style, text) runs; style in {plain, bold, italic, code}."""
    out = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append(("bold", part[2:-2]))
        elif part.startswith("*") and part.endswith("*"):
            out.append(("italic", part[1:-1]))
        elif part.startswith("`") and part.endswith("`"):
            out.append(("code", part[1:-1]))
        else:
            out.append(("plain", part))
    return out


def strip_inline(text: str) -> str:
    return re.sub(r"\*\*|\*|`", "", text)


# --------------------------------------------------------------------------- #
# Renderers
# --------------------------------------------------------------------------- #
def render_docx(blocks: list[Block], out: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.15

    def add_runs(par, text):
        for kind, chunk in split_inline(text):
            run = par.add_run(chunk)
            if kind == "bold":
                run.bold = True
            elif kind == "italic":
                run.italic = True
            elif kind == "code":
                run.font.name = "Consolas"
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    for b in blocks:
        if b.kind == "heading":
            doc.add_heading(strip_inline(b.text), level=min(b.level, 4))
        elif b.kind == "para":
            add_runs(doc.add_paragraph(), b.text)
        elif b.kind == "bullets":
            for item in b.items:
                add_runs(doc.add_paragraph(style="List Bullet"), item)
        elif b.kind == "numbers":
            for item in b.items:
                add_runs(doc.add_paragraph(style="List Number"), item)
        elif b.kind == "quote":
            par = doc.add_paragraph()
            par.paragraph_format.left_indent = Inches(0.35)
            add_runs(par, b.text)
            for run in par.runs:
                run.italic = True
        elif b.kind == "table" and b.rows:
            table = doc.add_table(rows=len(b.rows), cols=len(b.rows[0]))
            table.style = "Light Grid Accent 1"
            for r, row in enumerate(b.rows):
                for c, cell in enumerate(row[: len(b.rows[0])]):
                    par = table.cell(r, c).paragraphs[0]
                    add_runs(par, cell)
                    if r == 0:
                        for run in par.runs:
                            run.bold = True
        elif b.kind == "figure":
            path = ROOT / b.path
            if path.exists():
                doc.add_picture(str(path), width=Inches(6.2))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(strip_inline(b.caption))
            run.italic = True
            run.font.size = Pt(9.5)

    doc.save(out)
    return out


LATEX_ESCAPES = {"&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
                 "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}",
                 "^": r"\textasciicircum{}"}


def tex_escape(text: str) -> str:
    out = text.replace("\\", r"\textbackslash{}")
    for k, v in LATEX_ESCAPES.items():
        out = out.replace(k, v)
    return out


def tex_inline(text: str) -> str:
    parts = []
    for kind, chunk in split_inline(text):
        esc = tex_escape(chunk)
        parts.append({"bold": rf"\textbf{{{esc}}}", "italic": rf"\emph{{{esc}}}",
                      "code": rf"\texttt{{{esc}}}"}.get(kind, esc))
    return "".join(parts)


def render_tex(blocks: list[Block], out: Path) -> Path:
    # Starred (unnumbered) sectioning: the Markdown source already carries its
    # own section numbers, and letting LaTeX add a second set produces headings
    # like "0.1  1. Topic".
    sec = {1: "section*", 2: "section*", 3: "subsection*", 4: "subsubsection*"}
    body = []
    title = next((b.text for b in blocks if b.kind == "heading" and b.level == 1), "Proposal")
    started = False

    for b in blocks:
        if b.kind == "heading":
            if b.level == 1 and not started:
                started = True
                continue
            body.append(rf"\{sec.get(b.level, 'subsubsection*')}{{{tex_inline(b.text)}}}")
        elif b.kind == "para":
            body.append(tex_inline(b.text))
        elif b.kind in ("bullets", "numbers"):
            env = "itemize" if b.kind == "bullets" else "enumerate"
            body.append(rf"\begin{{{env}}}[leftmargin=1.3em,itemsep=2pt]")
            body.extend(rf"  \item {tex_inline(i)}" for i in b.items)
            body.append(rf"\end{{{env}}}")
        elif b.kind == "quote":
            body.append(rf"\begin{{quote}}\emph{{{tex_inline(b.text)}}}\end{{quote}}")
        elif b.kind == "table" and b.rows:
            ncol = len(b.rows[0])
            # First column left-aligned and allowed to wrap; the rest centred.
            spec = "p{0.29\\linewidth}" + "c" * (ncol - 1) if ncol > 1 else "l"
            body.append(r"\begin{center}\small")
            body.append(rf"\begin{{tabular}}{{{spec}}}\toprule")
            body.append(" & ".join(rf"\textbf{{{tex_inline(c)}}}" for c in b.rows[0]) + r" \\ \midrule")
            for row in b.rows[1:]:
                cells = (row + [""] * ncol)[:ncol]
                body.append(" & ".join(tex_inline(c) for c in cells) + r" \\")
            body.append(r"\bottomrule\end{tabular}")
            body.append(r"\end{center}")
        elif b.kind == "figure":
            path = ROOT / b.path
            if path.exists():
                body.append(r"\begin{figure}[htbp]\centering")
                # Relative to docs/ so the .tex compiles on any checkout (and Overleaf).
                rel = Path(os.path.relpath(path, DOCS)).as_posix()
                body.append(rf"\includegraphics[width=0.94\linewidth]{{{rel}}}")
                body.append(rf"\caption{{{tex_inline(b.caption)}}}")
                body.append(r"\end{figure}")
        body.append("")

    preamble = r"""\documentclass[11pt,a4paper]{article}
\usepackage[margin=2.3cm]{geometry}
\usepackage{fontspec}
\usepackage{booktabs}
\usepackage{graphicx}
\usepackage{enumitem}
\usepackage{microtype}
\usepackage[hidelinks]{hyperref}
\usepackage{caption}
\captionsetup{font=small,labelformat=empty,justification=raggedright,singlelinecheck=false}
\setmainfont{DejaVu Serif}
\setsansfont{DejaVu Sans}
\setmonofont{DejaVu Sans Mono}
\linespread{1.06}
\setlength{\parskip}{0.55em}
\setlength{\parindent}{0pt}
\title{%TITLE%}
\date{}
\begin{document}
\maketitle
""".replace("%TITLE%", tex_escape(title))

    out.write_text(preamble + "\n".join(body) + "\n\\end{document}\n")
    return out


def render_txt(blocks: list[Block], out: Path) -> Path:
    """Plain prose, matching the register of the original proposal draft."""
    lines: list[str] = []
    for b in blocks:
        if b.kind == "heading":
            text = strip_inline(b.text)
            if b.level == 1:
                lines += [text, "=" * len(text), ""]
            else:
                lines += ["", text, ""]
        elif b.kind == "para":
            lines += [strip_inline(b.text), ""]
        elif b.kind == "bullets":
            lines += [f"  - {strip_inline(i)}" for i in b.items] + [""]
        elif b.kind == "numbers":
            lines += [f"  {n}. {strip_inline(i)}" for n, i in enumerate(b.items, 1)] + [""]
        elif b.kind == "quote":
            lines += [f"  | {strip_inline(b.text)}", ""]
        elif b.kind == "table" and b.rows:
            widths = [max(len(strip_inline(r[c])) if c < len(r) else 0 for r in b.rows)
                      for c in range(len(b.rows[0]))]
            for n, row in enumerate(b.rows):
                cells = (row + [""] * len(widths))[: len(widths)]
                lines.append("  " + "  ".join(strip_inline(c).ljust(w)
                                              for c, w in zip(cells, widths)).rstrip())
                if n == 0:
                    lines.append("  " + "  ".join("-" * w for w in widths))
            lines.append("")
        elif b.kind == "figure":
            lines += [f"  [Figure: {strip_inline(b.caption)}  ({b.path})]", ""]
    return _write(out, "\n".join(lines))


def _write(path: Path, text: str) -> Path:
    path.write_text(text)
    return path


def build_pdf(tex: Path) -> Path | None:
    xelatex = str(XELATEX) if XELATEX.exists() else shutil.which("xelatex")
    if not xelatex:
        print(f"xelatex not found at {XELATEX} or on PATH; skipping PDF")
        return None
    for _ in range(2):                       # second pass resolves references
        proc = subprocess.run(
            [xelatex, "-interaction=nonstopmode", "-halt-on-error", tex.name],
            capture_output=True, text=True, cwd=tex.parent,
        )
    pdf = tex.with_suffix(".pdf")
    if not pdf.exists():
        tail = "\n".join(proc.stdout.strip().split("\n")[-25:])
        print(f"xelatex failed:\n{tail}")
        return None
    for ext in (".aux", ".log", ".out"):
        tex.with_suffix(ext).unlink(missing_ok=True)
    return pdf


def main() -> None:
    if not SRC.exists():
        raise SystemExit(f"missing {SRC}")
    md = SRC.read_text()
    validate(md)
    blocks = parse(md)
    print(f"parsed {len(blocks)} blocks from {SRC.relative_to(ROOT)}")

    print("wrote", render_docx(blocks, DOCS / "proposal.docx").relative_to(ROOT))
    print("wrote", render_txt(blocks, ROOT / "proposal.txt").relative_to(ROOT))
    tex = render_tex(blocks, DOCS / "proposal.tex")
    print("wrote", tex.relative_to(ROOT))
    pdf = build_pdf(tex)
    if pdf:
        print("wrote", pdf.relative_to(ROOT))


if __name__ == "__main__":
    sys.exit(main())
