"""EFSA/ECDC EU Summary Report on AMR: designed One Health monitoring.

Why this source matters more than its size suggests. Every other real dataset in
this project is *convenience-sampled*: NCBI Pathogen Detection contains whatever
laboratories chose to sequence and deposit, so its denominators reflect
submission behaviour rather than a sampling frame. The EU Summary Report is the
opposite. Member States are legally required to monitor specified
pathogen-matrix combinations, isolates are tested against a harmonised panel,
and resistance is called against common EUCAST epidemiological cut-off values.
Human and animal arms are reported for the *same* organisms, in annexes whose
sheet names state the sector outright.

That makes it the natural control for our central negative result. If reservoir
sectors fail to track the clinical sector in NCBI, the obvious objection is that
NCBI's sampling is too ragged to show it. EFSA answers that objection.

**What this data cannot do.** Only two EU Summary Report annex sets are
published in machine-readable form (the 2023 and 2024 reports, Zenodo records
14645440 and 17950222, CC-BY-4.0), and each sheet holds a single year. Two
annual observations cannot support a lead-lag analysis, and none is attempted
here. The "Temporal trends" sheets look promising and are not: they summarise
which Member States increased or decreased, not the underlying series.

So the test this module runs is **cross-sectional coupling** -- across Member
States, does animal resistance track human resistance for the same
organism-antimicrobial pair? That is a weaker claim than lead-lag but a
necessary condition for it, and it is the analysis the published literature
performs on these annexes. That the strongest One Health dataset in Europe
supports only a cross-sectional test is itself a finding, and it belongs
alongside the others in the report.
"""
from __future__ import annotations

import argparse
import json
import re
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd

from ..config import DATA_PROCESSED, DATA_RAW, RESULTS

#: Zenodo records holding machine-readable annexes, newest first. CC-BY-4.0.
ZENODO_RECORDS = {2024: "17950222", 2023: "14645440"}

EFSA_RAW = DATA_RAW / "efsa"

#: Annexes with a human arm and an animal arm for the same organism.
ANNEX_PATTERNS = {
    "human": re.compile(r"(Salmonella|Campylobacter)_humans", re.I),
    "animal": re.compile(r"(Salmonella_food_producing_animals|Campylobacter_(?:animals|\d{4}))", re.I),
}

#: The animal annexes use three-letter codes, the human annexes use full names.
#: Only the intersection can be compared across sectors.
ANTIMICROBIAL_CODES = {
    "GEN": "Gentamicin", "CHL": "Chloramphenicol", "CIP": "Ciprofloxacin",
    "ERY": "Erythromycin", "TET": "Tetracycline", "AMP": "Ampicillin",
    "FOT": "Cefotaxime", "CTX": "Cefotaxime", "TAZ": "Ceftazidime",
    "CAZ": "Ceftazidime", "MERO": "Meropenem", "MEM": "Meropenem",
    "NAL": "Nalidixic acid", "SMX": "Sulfamethoxazole", "TMP": "Trimethoprim",
    "AZI": "Azithromycin", "AZM": "Azithromycin", "COL": "Colistin",
    "TGC": "Tigecycline", "ETP": "Ertapenem", "AMK": "Amikacin",
}

#: Columns that look like antimicrobials but are not: complete susceptibility,
#: multidrug resistance, and combined-resistance summaries. Correlating these
#: across sectors would double-count the single-drug columns they derive from.
NOT_AN_ANTIMICROBIAL = re.compile(r"^(CS|MDR|Combined|Co[-_ ]?resistance)", re.I)

#: EFSA tables end with aggregate rows -- 'Total (23 MSs)', 'Median MSs',
#: 'Total non-MSs'. They are not countries. Left in, they add a point to every
#: correlation that is by construction the weighted mean of the others, which
#: inflates the coupling estimate this module exists to measure.
AGGREGATE_ROW = re.compile(r"^(total|median|mean|overall)\b|\bMSs?\b", re.I)

#: The human and animal annexes abbreviate serovars differently. Resistance in
#: broiler C. coli must be compared with human C. coli, not with a human average
#: that mixes serovars.
SEROVAR_ALIASES = {
    "S. Enter": "S. Enteritidis", "S. Typhim": "S. Typhimurium",
    "S. Infan": "S. Infantis", "S. Kenty": "S. Kentucky",
    "mono S. Typhim": "S. monophasic", "Salm spp.": "Salmonella spp.",
}


def _normalise_serovar(value: str) -> str:
    v = (value or "").strip()
    return SEROVAR_ALIASES.get(v, v)


# --------------------------------------------------------------------------- #
# Fetch
# --------------------------------------------------------------------------- #
def _curl_json(url: str) -> dict:
    """Fetch JSON with curl rather than requests.

    ``requests`` reliably times out against zenodo.org from this host while
    curl succeeds against the same URL -- a TLS/CDN negotiation quirk rather
    than anything about the request. Every other source in this project uses
    ``requests``; this is the one exception and it is deliberate.
    """
    out = subprocess.run(["curl", "-sL", "--max-time", "120", url],
                         capture_output=True, check=True)
    return json.loads(out.stdout)


def fetch_annexes(force: bool = False) -> list[Path]:
    """Download the human and animal annex spreadsheets, caching under data/raw."""
    EFSA_RAW.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    for year, record in ZENODO_RECORDS.items():
        meta = _curl_json(f"https://zenodo.org/api/records/{record}")
        licence = (meta.get("metadata", {}).get("license") or {}).get("id")
        for f in meta.get("files", []):
            key = f.get("key", "")
            if not key.lower().endswith((".xlsx", ".xlsm")):
                continue
            if not any(p.search(key) for p in ANNEX_PATTERNS.values()):
                continue
            dest = EFSA_RAW / f"{year}__{key.replace(' ', '_')}"
            if dest.exists() and not force and dest.stat().st_size > 1000:
                paths.append(dest)
                continue
            subprocess.run(["curl", "-sL", "--max-time", "300",
                            f["links"]["self"], "-o", str(dest)], check=True)
            print(f"  {dest.stat().st_size/1e3:8.1f} KB  {dest.name}  [{licence}]")
            paths.append(dest)
    return paths


# --------------------------------------------------------------------------- #
# Parsing
# --------------------------------------------------------------------------- #
def _clean_country(value) -> str | None:
    """Strip EFSA footnote markers: 'Bulgaria(b)' -> 'Bulgaria'."""
    if not isinstance(value, str):
        return None
    name = re.sub(r"\s*\([a-z]\)\s*$", "", value.strip())
    if not name or AGGREGATE_ROW.search(name):
        return None
    return name


def _num(value) -> float:
    """Parse a cell that may be '-', 'NA', '<0.1' or a number."""
    if value is None:
        return np.nan
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    if s in {"", "-", "–", "NA", "N/A", "."}:
        return np.nan
    s = s.replace("<", "").replace(">", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return np.nan


def parse_human_sheet(ws, sheet: str) -> pd.DataFrame:
    """Human annexes: row 2 holds antimicrobial names, row 3 holds N / % Res pairs."""
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    if len(rows) < 4:
        return pd.DataFrame()
    drugs, labels = rows[1], rows[2]

    # Each antimicrobial spans two columns; the name sits above the 'N' column.
    columns: list[tuple[int, str]] = []
    current = None
    for i, (d, lab) in enumerate(zip(drugs, labels)):
        if isinstance(d, str) and d.strip():
            current = d.strip()
        if isinstance(lab, str) and lab.strip().lower().startswith("% res") and current:
            if not NOT_AN_ANTIMICROBIAL.match(current):
                columns.append((i, current))

    out = []
    for r in rows[3:]:
        country = _clean_country(r[0])
        if not country or country.lower().startswith(("n:", "%", "note", "(", "–", "-")):
            continue
        for idx, drug in columns:
            if idx >= len(r):
                continue
            n = _num(r[idx - 1])
            pct = _num(r[idx])
            if np.isnan(pct):
                continue
            out.append({"country": country, "antimicrobial": drug,
                        "n": n, "percent_resistant": pct})
    return pd.DataFrame(out)


def parse_animal_sheet(ws, sheet: str) -> pd.DataFrame:
    """Animal annexes: a single header row of 'N' plus '<CODE> (%)' columns."""
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    if len(rows) < 2:
        return pd.DataFrame()
    header = rows[0]
    if not (isinstance(header[0], str) and "country" in header[0].lower()):
        return pd.DataFrame()

    n_idx = next((i for i, h in enumerate(header)
                  if isinstance(h, str) and h.strip().upper() == "N"), None)
    drug_cols = {}
    for i, h in enumerate(header):
        if not isinstance(h, str):
            continue
        m = re.match(r"^([A-Z]{2,4})\s*\(%\)$", h.strip())
        if m and not NOT_AN_ANTIMICROBIAL.match(m.group(1)):
            drug_cols[i] = ANTIMICROBIAL_CODES.get(m.group(1), m.group(1))

    out = []
    for r in rows[1:]:
        country = _clean_country(r[0])
        if not country:
            continue
        n = _num(r[n_idx]) if n_idx is not None else np.nan
        for idx, drug in drug_cols.items():
            if idx >= len(r):
                continue
            pct = _num(r[idx])
            if np.isnan(pct):
                continue
            out.append({"country": country, "antimicrobial": drug,
                        "n": n, "percent_resistant": pct})
    return pd.DataFrame(out)


#: Sheet-name prefixes that identify the animal population. Sheets whose name
#: contains 'proportion', 'prevalence' or 'trend' report something other than
#: resistance occurrence and are skipped.
ANIMAL_POPULATIONS = ("Broilers", "Turkeys", "Pigs", "Cattle", "Calves", "Laying hens")
SKIP_SHEETS = re.compile(r"proportion|prevalence|trend|MDR|CS_|contents|patterns|levels", re.I)


def parse_annex(path: Path) -> pd.DataFrame:
    """Parse one annex workbook into a long panel."""
    import openpyxl

    report_year = int(path.name.split("__")[0])
    is_human = bool(ANNEX_PATTERNS["human"].search(path.name))
    organism = "Salmonella" if re.search(r"Salmonella", path.name, re.I) else "Campylobacter"

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    frames = []
    for sheet in wb.sheetnames:
        if SKIP_SHEETS.search(sheet):
            continue
        ws = wb[sheet]
        if is_human:
            if not sheet.lower().startswith("hum"):
                continue
            df = parse_human_sheet(ws, sheet)
            population = "humans"
            serovar = _normalise_serovar(re.sub(r"^Hum\s*", "", sheet))
        else:
            population = next((p for p in ANIMAL_POPULATIONS if sheet.startswith(p)), None)
            if population is None:
                continue
            df = parse_animal_sheet(ws, sheet)
            serovar = _normalise_serovar(sheet[len(population):])
        if df.empty:
            continue
        df["report_year"] = report_year
        df["organism"] = organism
        df["sector"] = "human" if is_human else "animal"
        df["population"] = population
        df["serovar"] = serovar
        df["sheet"] = sheet
        frames.append(df)
    wb.close()
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def build_panel(force: bool = False) -> pd.DataFrame:
    """Parse every downloaded annex into one long panel."""
    paths = fetch_annexes(force=force)
    frames = [parse_annex(p) for p in sorted(paths)]
    frames = [f for f in frames if not f.empty]
    panel = pd.concat(frames, ignore_index=True)
    # Denominators below 10 are reported by EFSA as 'NA' but some slip through;
    # a percentage on five isolates is noise whatever the sampling frame.
    return panel[(panel["n"].isna()) | (panel["n"] >= 10)].reset_index(drop=True)


# --------------------------------------------------------------------------- #
# Cross-sectional coupling
# --------------------------------------------------------------------------- #
def sector_coupling(panel: pd.DataFrame, min_countries: int = 8) -> dict:
    """Across Member States, does animal resistance track human resistance?

    For each organism-antimicrobial pair, correlate the animal and human
    resistance percentages across countries. Spearman rather than Pearson: the
    percentages are bounded, skewed, and a handful of countries sit at 0 or 100.

    This is a necessary condition for the lead-lag structure our early-warning
    system assumes. It is not sufficient -- two sectors can be coupled without
    either leading -- and the two available years cannot distinguish those.
    """
    from scipy.stats import spearmanr

    occurrence = panel[panel["sector"].isin(["human", "animal"])]
    findings = []
    # Grouped by serovar as well as organism and antimicrobial: comparing broiler
    # C. coli against a human average over C. coli and C. jejuni would be
    # measuring two different bugs against each other.
    for (organism, serovar, drug), block in occurrence.groupby(
            ["organism", "serovar", "antimicrobial"]):
        human = (block[block["sector"] == "human"]
                 .groupby("country")["percent_resistant"].mean())
        for population, animal_block in block[block["sector"] == "animal"].groupby("population"):
            animal = animal_block.groupby("country")["percent_resistant"].mean()
            shared = human.index.intersection(animal.index)
            if len(shared) < min_countries:
                continue
            h, a = human.loc[shared].to_numpy(), animal.loc[shared].to_numpy()
            if np.std(h) == 0 or np.std(a) == 0:
                continue
            rho, p = spearmanr(a, h)
            findings.append({
                "organism": organism, "serovar": serovar,
                "antimicrobial": drug, "population": population,
                "n_countries": int(len(shared)), "spearman_rho": round(float(rho), 3),
                "p_value": float(p),
            })

    df = pd.DataFrame(findings)
    if df.empty:
        return {"n_pairs": 0, "findings": []}

    # Are the correlations positive more often than chance? A sign test over
    # pairs is the honest summary: individual correlations on 10-25 countries
    # are noisy, but their direction in aggregate is informative.
    from scipy.stats import binomtest

    positive = int((df["spearman_rho"] > 0).sum())
    sign_p = float(binomtest(positive, len(df), 0.5, alternative="greater").pvalue)
    return {
        "n_pairs": int(len(df)),
        "n_positive": positive,
        "median_rho": float(df["spearman_rho"].median()),
        "sign_test_p": sign_p,
        "n_significant_positive": int(((df["p_value"] < 0.05) & (df["spearman_rho"] > 0)).sum()),
        "countries": sorted(panel["country"].unique().tolist()),
        "findings": df.sort_values("spearman_rho", ascending=False).to_dict("records"),
        "resolution_caveat": (
            "Only two EU Summary Report annex sets are published in machine-readable "
            "form, each holding a single year, so no lead-lag analysis is possible on "
            "this source. This is a cross-sectional coupling test across Member States: "
            "a necessary but not sufficient condition for the temporal structure an "
            "early-warning system needs."
        ),
    }


# --------------------------------------------------------------------------- #
# CLI
# --------------------------------------------------------------------------- #
def main() -> None:
    ap = argparse.ArgumentParser(description="EFSA/ECDC EU Summary Report on AMR.")
    ap.add_argument("--force", action="store_true", help="re-download annexes")
    args = ap.parse_args()

    panel = build_panel(force=args.force)
    panel.to_csv(DATA_PROCESSED / "efsa_panel.csv", index=False)
    print(f"\npanel: {len(panel):,} rows | {panel['country'].nunique()} countries | "
          f"{panel['organism'].nunique()} organisms | "
          f"{panel['antimicrobial'].nunique()} antimicrobials")
    print(panel.groupby(["sector", "population"]).size().to_string())

    study = sector_coupling(panel)
    (RESULTS / "efsa_coupling_study.json").write_text(json.dumps(study, indent=2, default=float))

    if not study["n_pairs"]:
        print("\nno comparable organism-antimicrobial pairs found")
        return
    print(f"\n{study['n_pairs']} organism x antimicrobial x population pairs across "
          f"{len(study['countries'])} countries")
    print(f"positive animal-human correlation in {study['n_positive']}/{study['n_pairs']} "
          f"pairs (sign test p = {study['sign_test_p']:.4g}); "
          f"median rho = {study['median_rho']:+.3f}; "
          f"{study['n_significant_positive']} individually significant")
    print("\nstrongest couplings:")
    for f in study["findings"][:10]:
        print(f"  {f['serovar']:<16s} {f['antimicrobial']:<16s} {f['population']:<12s} "
              f"rho={f['spearman_rho']:+.2f} p={f['p_value']:.3g} n={f['n_countries']}")
    print(f"\n{study['resolution_caveat']}")


if __name__ == "__main__":
    main()
